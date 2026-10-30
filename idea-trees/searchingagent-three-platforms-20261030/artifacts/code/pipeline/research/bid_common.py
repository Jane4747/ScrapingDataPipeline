"""医卫数字化中标采集共用的口径与出表逻辑。

央采（`pipeline/research/zycg.py`）与北京/雄安（`pipeline/research/bjxa.py`）共用同一套
八列口径：谁算医卫主体、什么算数字化标的、场景怎么归、金额怎么折万元、中标单位怎么判运营商。
把口径放在一处，避免两个采集器各自解释一遍，最后两张表对不上。

本模块只做纯函数，不联网；站点差异（列表怎么翻页、正文长什么样）留在各自的采集器里。
"""
from __future__ import annotations

import csv
import re
from pathlib import Path

COLUMNS = ["项目名称", "产品", "招标单位", "客户标签", "中标单位", "运营商类型", "中标金额（万元）", "场景"]

MED_ORG = re.compile(
    r"医院|卫生院|社区卫生|疾病预防控制|疾控中心|卫生健康委|卫生局|卫健委|中医药管理|医疗保障局|"
    r"医保|妇幼保健|血液中心|血站|急救中心|康复中心|医学科学院|医学研究所|生物医学|公共卫生|"
    r"医疗|医学|护理|药学|肿瘤|口腔|眼科|精神卫生|妇产|儿童医院|心血管|胸科|传染病|骨科|皮肤"
)
DIGITAL = re.compile(
    r"信息化|信息系统|信息平台|信息设备|信息交换|网络|服务器|存储|备份|容灾|数据库|数据治理|大数据|"
    r"数据平台|算力|云计算|云平台|云节点|超融合|虚拟化|机房|数据中心|安全|等保|密码|防火墙|"
    r"堡垒机|态势感知|日志审计|加密|负载均衡|交换机|路由|无线|链路|专线|系统|平台|软件|应用|"
    r"门户|中台|接口|集成|运维|维保|终端|计算机|打印|桌面|会议|监控|智慧|数字化|智能|人工智能|"
    r"大模型|AI|PACS|LIS|HIS|EMR|电子病历|影像|检验|病理|超声|内镜|心电|远程|会诊|互联网|"
    r"科研|DRG|DIP|病案|无纸化|电子签名|自助|排队叫号|随访问卷|"
    r"数智|硬件|软硬件|数据|会商|指挥中心|指挥调度|机房设备|扩容"
)
# 纯设备/工程/后勤类：命中则判定为“非数字化”，即便同时出现医卫与系统类词
NON_DIGITAL = re.compile(
    r"餐饮|食堂|保洁|物业|安保|保安|绿化|被服|洗涤|布草|停车|电梯|班车|印刷|家具|办公用品|"
    r"试剂|耗材|药品|疫苗|手术器械|消毒|供氧|氧气|液氧|被褥|工作服|体检服务|医疗废物|污水|"
    r"劳务|人力|保险|租赁|审计|咨询|翻译|设计服务|改造工程|装修|土建|防水|消防工程|"
    r"施工|综合维修|配电运维|监理|勘察|造价|保洁服务|绿化养护"
)

STAGE_WORDS = [
    ("云与算力", r"云|算力|虚拟化|超融合|容器|GPU|服务器|资源池|节点扩容"),
    # “大数据中心/数据中心”常是机构名，只有明确机房要素才归数据中心与机房
    ("数据中心与机房", r"机房|数据中心机房|UPS|精密空调|机柜|综合布线|灾备中心|机房建设"),
    # “网络安全”属信息安全，必须先于“网络与链路”判定
    ("信息安全", r"安全|等保|密码|防火墙|堡垒机|态势感知|日志审计|加密|防病毒|入侵|零信任|签名"),
    ("网络与链路", r"网络|交换机|路由|无线|链路|专线|负载均衡|带宽|流量"),
    ("存储与备份", r"存储|备份|容灾|归档|磁带|磁盘阵列|NAS|SAN"),
    ("影像与检验", r"PACS|影像|放射|超声|内镜|检验|LIS|病理|心电|核医学|内窥镜"),
    ("电子病历与HIS", r"HIS|电子病历|EMR|医嘱|临床路径|DRG|DIP|病案|手麻|手术麻醉|合理用药"),
    ("互联网医院与远程医疗", r"互联网医院|互联网诊疗|远程|会诊|在线问诊|随访"),
    ("人工智能", r"人工智能|AI|大模型|智能辅助|辅助诊断|智能审方|语音"),
    ("应用系统与平台", r"系统|平台|软件|应用|门户|中台|科研|数据|档案|排队叫号|自助"),
    ("运维与技术服务", r"运维|维保|服务外包|技术支持|驻场|租赁服务"),
    ("终端与办公设备", r"计算机|终端|打印|桌面|显示|会议|投影|扫描"),
]


def clean(text: str) -> str:
    return re.sub(r"\s+", " ", (text or "").replace("\u3000", " ")).strip()


def classify_scope(title: str, buyer: str = "") -> tuple[bool, str]:
    """判断是否属于“医卫数字化”。返回 (是否收录, 原因)。"""
    haystack = f"{title} {buyer}"
    medical = bool(MED_ORG.search(haystack))
    digital = bool(DIGITAL.search(title))
    blocked = NON_DIGITAL.search(title)
    if not medical:
        return False, "非医卫采购主体"
    if blocked:
        return False, "医卫主体的非数字化采购（后勤/设备/工程类）"
    if not digital:
        return False, "医卫采购但标题未见数字化标的，需人工复核"
    return True, ""


def customer_tag(buyer: str, context: str = "") -> str:
    if re.search(r"疾病预防控制|疾控中心", buyer):
        return "疾控机构"
    if re.search(r"医疗保障局|医保中心|医保局", buyer):
        return "医保机构"
    if re.search(r"卫生健康委|卫生健康局|卫生局|中医药管理|药品监督管理", buyer):
        return "卫生行政/监管"
    if re.search(r"妇幼保健", buyer):
        return "妇幼保健机构"
    if re.search(r"考试中心|统计信息中心|医疗管理服务指导中心|妇幼健康中心|医院管理中心|"
                 r"卫生发展研究中心|医学情报|人才交流|药具管理|卫生监督", buyer):
        return "卫生支撑机构"
    if re.search(r"血液中心|血站|急救中心|紧急救援", buyer):
        return "公共卫生机构（血液/急救）"
    if re.search(r"中国医学科学院|中国疾病预防控制中心|中国中医科学院", buyer):
        return "国家级医学科研/医疗机构"
    if re.search(r"军|武警|战区|部队", buyer):
        return "军队医疗机构"
    if re.search(r"大学|学院|学校", buyer) and re.search(r"医院|附属", buyer):
        return "高校附属医院"
    if re.search(r"大学|学院|学校", buyer) and re.search(r"医学|医|卫生", buyer + context):
        return "医学教育/科研机构"
    if re.search(r"中医", buyer):
        return "中医医疗机构"
    if re.search(r"肿瘤|口腔|眼科|精神|妇产|儿童医院|心血管|胸科|传染|血液病|康复|骨科|皮肤|"
                 r"耳鼻|咽喉|整形|职业病|老年|肛肠|结核|麻风|地方病|口腔|烧伤|创伤", buyer):
        return "专科医疗机构"
    if re.search(r"社区卫生|卫生院|基层", buyer):
        return "基层医疗机构"
    if re.search(r"医院", buyer):
        return "综合医疗机构"
    return "其他医卫机构"


def carrier_type(supplier: str) -> str:
    if re.search(r"中国移动|中移|移动通信", supplier):
        return "中国移动"
    if re.search(r"中国电信|中电信|电信集团|电信股份", supplier):
        return "中国电信"
    if re.search(r"中国联通|中联通|联通", supplier):
        return "中国联通"
    if re.search(r"中国广电|广电网络", supplier):
        return "中国广电"
    return "非运营商"


def scene_of(text: str) -> str:
    for name, pattern in STAGE_WORDS:
        if re.search(pattern, text, re.I):
            return name
    return "其他数字化"


# 口径边界：医卫主体 + 数字化词命中，但标的其实是硬件维保/设备购置/后勤设施。
# 这类不该算数字化采购，可收紧口径会让央采与北京两张表同时变动，故先标出来供人工复核。
BOUNDARY_RULES = [
    ("医工设备维保", r"(医用|医疗|放射|核磁|CT|DR|超声|内镜|直线加速器|加速器|显微镜|监护|DSA|"
                     r"血管机|胃肠机|钼靶|骨密度).{0,10}(维保|保修|维修|保养)"),
    ("医用设备购置", r"带量采购|医用设备集中|医疗设备采购|医疗设备购置|科研设备购置|仪器公共平台|"
                     r"(?:医疗|医用)设备(?!管理|信息|系统|平台|数据|软件|网络|运维|监控)"),
    ("检验病理委外", r"(委外|外送|委托).{0,6}(检测|检验|病理)|(检测|检验).{0,4}委外"),
    ("后勤设施维保", r"(空调|净化|电梯|锅炉|污水处理|物流|被服|食堂).{0,12}(维保|运行|服务|采购)"),
    ("非数字化物资", r"农副产品|食材配送|副食品"),
]


def boundary_hint(title: str) -> str:
    """返回口径边界类别，空串表示未被标记。"""
    for name, pattern in BOUNDARY_RULES:
        if re.search(pattern, title or ""):
            return name
    return ""


def amount_wan(text: str) -> float | None:
    """金额统一折算为万元；带“元”的按 /10000。"""
    if not text:
        return None
    body = text.replace(",", "").replace("，", "").replace(" ", "")
    match = re.search(r"人民币\s*([0-9]+(?:\.[0-9]+)?)\s*(万元|元|亿元)?", body)
    if not match:
        match = re.search(r"([0-9]+(?:\.[0-9]+)?)\s*(万元|元|亿元)", body)
    if not match:
        return None
    value = float(match.group(1))
    unit = match.group(2) or "元"
    if unit == "亿元":
        return round(value * 10000, 4)
    if unit == "元":
        return round(value / 10000, 4)
    return round(value, 4)


def amount_field(text: str, label: str) -> str:
    """金额常写成“成交金额：人民币\\n697000.00元”，要跨行取到单位为止。"""
    match = re.search(rf"{label}[：:]?\s*([\s\S]{{0,80}})", text)
    if not match:
        return ""
    segment = match.group(1)
    unit = re.search(r"(万元|亿元|元)", segment)
    if unit:
        segment = segment[:unit.end()]
    return clean(segment)


def labeled_lines(text: str) -> dict[str, str]:
    """把“1．项目编号： xxx”这类行拆成标签字典，同名标签保留首个非空值。"""
    fields: dict[str, str] = {}
    for raw in text.split("\n"):
        line = clean(raw)
        if not line or len(line) > 400:
            continue
        match = re.match(r"^(?:\d+)?[．.、]?\s*([\u4e00-\u9fa5（）()]{2,16})[：:]\s*(.+)$", line)
        if not match:
            continue
        key, value = match.group(1), match.group(2).strip()
        if not value:
            continue
        fields.setdefault(key, value)
    return fields


def field(text: str, label: str, pattern: str = r"[^\n]{1,60}") -> str:
    """标签与值之间可能被换行拆开，先按整段文本找，再退化到逐行找。"""
    match = re.search(rf"{label}[：:]\s*({pattern})", text)
    if match:
        return clean(match.group(1))
    match = re.search(rf"{label}[：:]\s*\n\s*({pattern})", text)
    return clean(match.group(1)) if match else ""


def build_row(detail: dict) -> dict:
    supplier = detail.get("中标单位", "")
    title = detail.get("公告标题", "") or detail.get("项目名称", "")
    goods = f"{detail.get('产品','')} {detail.get('项目名称','')}"
    # 场景以标题为准：标题说“数据治理平台”就不该被标的里的“存储”带到存储类
    scene = scene_of(title)
    if scene == "其他数字化":
        scene = scene_of(goods)
    return {
        "项目名称": detail.get("项目名称", ""),
        "产品": detail.get("产品", "") or scene,
        "招标单位": detail.get("招标单位", ""),
        "客户标签": customer_tag(detail.get("招标单位", ""), title),
        "中标单位": supplier,
        "运营商类型": carrier_type(supplier) if supplier else "",
        "中标金额（万元）": detail.get("中标金额（万元）"),
        "场景": scene,
    }


def write_csv(rows: list[dict], path: Path, columns: list[str] | None = None) -> None:
    columns = columns or COLUMNS
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=columns)
        writer.writeheader()
        for row in rows:
            writer.writerow({k: ("" if row.get(k) is None else row.get(k)) for k in columns})


def write_bid_workbook(records: list[dict], out_dir: Path, stem: str, sheet_title: str,
                       detail_columns: list[str]) -> list[Path]:
    """主表固定八列，明细页补编号/预算/链接等复核字段。缺 openpyxl 时只落 CSV。"""
    out_dir.mkdir(parents=True, exist_ok=True)
    rows = [r["row"] for r in records if r.get("row")]
    csv_path = out_dir / f"{stem}.csv"
    write_csv(rows, csv_path)
    written = [csv_path]

    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter
    except ImportError:
        return written

    book = Workbook()
    sheet = book.active
    sheet.title = sheet_title
    sheet.append(COLUMNS)
    header_fill = PatternFill("solid", fgColor="1F4E78")
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
    for row in rows:
        sheet.append([row.get(k) for k in COLUMNS])
    for index, width in enumerate([46, 30, 34, 18, 34, 12, 16, 18], 1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    for line in sheet.iter_rows(min_row=2):
        for cell in line:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:{get_column_letter(len(COLUMNS))}{sheet.max_row}"

    detail = book.create_sheet("明细")
    detail.append(["项目名称", "招标单位", "中标单位"] + detail_columns)
    for record in records:
        d = record.get("detail") or {}
        detail.append([d.get("项目名称", ""), d.get("招标单位", ""), d.get("中标单位", "")] +
                      [d.get(k, "") for k in detail_columns])
    detail.column_dimensions["A"].width = 46
    for index in range(len(detail_columns) + 1, len(detail_columns) + 5):
        detail.column_dimensions[get_column_letter(index)].width = 24
    detail.freeze_panes = "A2"
    xlsx_path = out_dir / f"{stem}.xlsx"
    book.save(xlsx_path)
    written.append(xlsx_path)
    return written
